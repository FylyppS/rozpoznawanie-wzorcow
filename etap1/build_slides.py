#!/usr/bin/env python3
"""Buduje slajdy etapu 1 na formatce wydziału (python-pptx).

    python etap1/build_slides.py --source-template ~/Downloads/preznainz.pptx

Z pełnej prezentacji-źródła zostaje tylko wzorzec z układami (zapis do etap1/formatka-wi.pptx,
bez slajdów i multimediów); na nim składane są slajdy. Treść pochodzi z utrwalonego protokołu
(cvbias/protocol.py), z surowych wyników próby uruchomienia (results/pilot) i z rysunków
(etap1/figures). Wzory są renderowane z notacji TeX (matplotlib mathtext) do PNG i wyśrodkowane.

Zakres slajdów = wymagania etapu 1 z karty tematu i kontraktu (ok. 10 minut). Do każdego slajdu
dołączone są notatki prelegenta (co powiedzieć, co policzyć na głos, na jakie pytanie się przygotować).
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import pandas as pd  # noqa: E402
from lxml import etree  # noqa: E402
from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from cvbias import protocol as P  # noqa: E402

# --- paleta: tekst w odcieniach szarości, akcent granatowy, karty jasnoszare/jasnoniebieskie -----
FONT = "Arial"
INK = RGBColor(0x0B, 0x0B, 0x0B)
INK2 = RGBColor(0x52, 0x51, 0x4E)
MUTED = RGBColor(0x89, 0x87, 0x81)
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
CARD = RGBColor(0xF3, 0xF5, 0xF8)
CARD_NAVY = RGBColor(0xE8, 0xEE, 0xF6)
ARROW = RGBColor(0xC3, 0xC2, 0xB7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ROW_ALT = RGBColor(0xF7, 0xF8, 0xFA)

# Obszar treści na slajdzie 13,333 × 7,5 cala: logotypy u góry, linia stopki na ~6,4 cala.
L, R_EDGE, TOP, BOTTOM = 0.92, 12.42, 1.85, 6.25
W = R_EDGE - L


# ---------------------------------------------------------------- pomocnicze: tekst i kształty
def set_bullet(paragraph, level: int, bullet: bool, indent: float = 0.24) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buFont"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    if bullet:
        pPr.set("marL", str(int(Inches(indent * (level + 1)))))
        pPr.set("indent", str(-int(Inches(indent))))
        bu_font = etree.SubElement(pPr, qn("a:buFont"))
        bu_font.set("typeface", "Arial")
        bu_char = etree.SubElement(pPr, qn("a:buChar"))
        bu_char.set("char", "•" if level == 0 else "–")
    else:
        pPr.set("marL", str(int(Inches(indent * level))))
        pPr.set("indent", "0")
        etree.SubElement(pPr, qn("a:buNone"))


def add_text(slide, x, y, w, h, items, size=16, color=INK, anchor="top", line_spacing=1.08,
             space_after=6, align="left"):
    """Pole tekstowe. ``items``: napisy albo słowniki {text|runs, level, bullet, bold, color, size, align,
    space_before, space_after}; ``runs`` to lista (tekst, opcje)."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
    for i, item in enumerate(items):
        if isinstance(item, str):
            item = {"text": item}
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = line_spacing
        p.space_before = Pt(item.get("space_before", 0))
        p.space_after = Pt(item.get("space_after", space_after))
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[item.get("align", align)]
        for text, opts in item.get("runs") or [(item.get("text", ""), {})]:
            run = p.add_run()
            run.text = text
            font = run.font
            font.name = FONT
            font.size = Pt(opts.get("size", item.get("size", size)))
            font.bold = opts.get("bold", item.get("bold", False))
            font.italic = opts.get("italic", item.get("italic", False))
            font.color.rgb = opts.get("color", item.get("color", color))
        set_bullet(p, item.get("level", 0), item.get("bullet", False))
    return box


def add_title(slide, text: str, size: int | None = None) -> None:
    """Tytuł w jednej linii pod logotypami; dłuższe tytuły dostają mniejszy stopień pisma."""
    if size is None:
        size = 30 if len(text) <= 48 else 26
    title = slide.shapes.title
    title.left, title.top, title.width, title.height = Inches(L), Inches(0.8), Inches(W), Inches(0.9)
    tf = title.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.name, run.font.size, run.font.bold = FONT, Pt(size), False
    run.font.color.rgb = INK


def add_card(slide, x, y, w, h, fill=CARD):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.adjustments[0] = 0.07
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def add_arrow(slide, x, y, w, h, direction="right"):
    kind = MSO_SHAPE.RIGHT_ARROW if direction == "right" else MSO_SHAPE.DOWN_ARROW
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = ARROW
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def add_table(slide, data, x, y, w, col_widths, row_h=0.36, size=12, align=None, bold_last=False):
    """Tabela z granatowym nagłówkiem i naprzemiennym tłem wierszy; ``align`` to lista wyrównań kolumn."""
    rows, cols = len(data), len(data[0])
    shape = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(row_h * rows))
    table = shape.table
    table.first_row, table.horz_banding = True, False
    for j, cw in enumerate(col_widths):
        table.columns[j].width = Inches(cw)
    for i in range(rows):
        table.rows[i].height = Inches(row_h)
        for j in range(cols):
            cell = table.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if i == 0 else (WHITE if i % 2 else ROW_ALT)
            p = cell.text_frame.paragraphs[0]
            p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[(align or ["l"] * cols)[j]]
            run = p.add_run()
            run.text = str(data[i][j])
            run.font.name, run.font.size = FONT, Pt(size)
            run.font.bold = i == 0 or (bold_last and i == rows - 1)
            run.font.color.rgb = WHITE if i == 0 else INK
    return shape


def formula_png(tex: str, path: Path, fontsize: int = 30, color: str = "#0b0b0b") -> Path:
    """Renderuje wzór TeX do przezroczystego PNG (matplotlib mathtext, czcionka Computer Modern)."""
    plt.rcParams["mathtext.fontset"] = "cm"
    fig = plt.figure(figsize=(0.1, 0.1), dpi=300)
    fig.patch.set_alpha(0)
    fig.text(0, 0, f"${tex}$", fontsize=fontsize, color=color)
    fig.savefig(path, dpi=300, bbox_inches="tight", pad_inches=0.04, transparent=True)
    plt.close(fig)
    return path


def add_formula(slide, tex: str, x, y, h, build_dir: Path, name: str, fontsize: int = 30,
                center_in: float | None = None, color: str = "#0b0b0b", max_w: float | None = None):
    """Wstawia wzór jako obraz o wysokości ``h`` cala. Z ``center_in`` = szerokość pola (w calach)
    obraz jest wyśrodkowany poziomo w polu zaczynającym się w ``x``; ``max_w`` (domyślnie = ``center_in``)
    ogranicza szerokość — za szeroki wzór jest proporcjonalnie pomniejszany."""
    path = formula_png(tex, build_dir / f"{name}.png", fontsize, color)
    with Image.open(path) as img:
        ratio = img.width / img.height
    w = h * ratio
    limit = max_w if max_w is not None else center_in
    if limit is not None and w > limit:
        w, h = limit, limit / ratio
    if center_in is not None:
        x = x + (center_in - w) / 2
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), height=Inches(h))


def add_notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


# ---------------------------------------------------------------- formatka
def make_blank_template(source: Path, target: Path) -> None:
    """Zostawia z prezentacji-źródła tylko wzorzec i układy (slajdy i multimedia znikają przy zapisie)."""
    prs = Presentation(str(source))
    sld_id_lst = prs.slides._sldIdLst
    for sld_id in list(sld_id_lst):
        prs.part.drop_rel(sld_id.rId)
        sld_id_lst.remove(sld_id)
    prs.save(str(target))


def pct(v: float) -> str:
    return f"{100 * v:.1f} %".replace(".", ",")


def num(v: float, nd: int = 3) -> str:
    return f"{v:.{nd}f}".replace(".", ",")


def tex_num(v: float, nd: int = 3) -> str:
    """Liczba do wzoru TeX z przecinkiem dziesiętnym (bez odstępu po przecinku)."""
    return f"{v:.{nd}f}".replace(".", "{,}")


# ---------------------------------------------------------------- slajdy
def build(template: Path, out: Path, results: Path, figures: Path, build_dir: Path) -> None:
    prs = Presentation(str(template))
    lay = {layout.name: layout for layout in prs.slide_layouts}
    folds = pd.read_csv(results / "reproduction_folds.csv")
    summary = pd.read_csv(results / "reproduction_summary.csv")
    own_folds = pd.read_csv(results / "own_folds.csv")
    meta = json.loads((results / "meta.json").read_text(encoding="utf-8"))
    fits = P.planned_fits()
    gen, gen_imb = P.GENERATORS["balanced"], P.GENERATORS["imbalanced"]

    def new_slide(title: str):
        slide = prs.slides.add_slide(lay["Tylko tytuł"])
        add_title(slide, title)
        return slide

    # 1. Tytuł ------------------------------------------------------------------------------
    s = prs.slides.add_slide(lay["Slajd tytułowy"])
    t = s.shapes.title
    t.left, t.top, t.width, t.height = Inches(1.2), Inches(1.5), Inches(10.9), Inches(2.0)
    p = t.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Obciążenie i zmienność oceny\nprzez walidację krzyżową"
    r.font.name, r.font.size, r.font.color.rgb = FONT, Pt(44), INK
    sub = s.placeholders[1]
    sub.left, sub.top, sub.width, sub.height = Inches(1.2), Inches(3.75), Inches(10.9), Inches(2.4)
    tf = sub.text_frame
    lines = [
        ("Etap 1 — prezentacja tematu  ·  temat 06, obszar: metodologia eksperymentów", 18, INK2),
        ("R. Kohavi, A Study of Cross-Validation and Bootstrap for Accuracy Estimation and Model Selection, IJCAI 1995", 15, MUTED),
        ("", 10, MUTED),
        ("Dariusz Salzburg (103931)  ·  Filip Sejud (103928)", 20, INK),
        ("Rozpoznawanie wzorców, studia II stopnia, 2026/2027, grupa 2  ·  29.10.2026", 15, INK2),
    ]
    for i, (text, size, color) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text
        r.font.name, r.font.size, r.font.color.rgb = FONT, Pt(size), color
        set_bullet(p, 0, False)
    add_notes(s, (
        "CO POWIEDZIEĆ (30 s): Temat 06 z obszaru „metodologia eksperymentów”. Publikacja: Kohavi, IJCAI 1995 — "
        "jedna z najczęściej cytowanych prac o walidacji krzyżowej. Pytanie projektu: jak liczba foldów k wpływa na "
        "obciążenie i zmienność błędu szacowanego przez CV i co zmienia stratyfikacja przy klasie rzadkiej. "
        "Plan prezentacji: publikacja → pytanie i hipotezy → teoria z przykładem ręcznym → plan badania → wykonalność."
    ))

    # 2. Publikacja i problem ----------------------------------------------------------------
    s = new_slide("Publikacja i problem")
    add_card(s, L, TOP, 6.0, 3.9)
    add_text(s, L + 0.2, TOP + 0.15, 5.6, 3.6, [
        {"text": "Kohavi (IJCAI 1995)", "bold": True, "size": 17, "color": NAVY, "space_after": 10},
        {"text": "Jak dobrze CV i bootstrap szacują trafność klasyfikatora?", "bullet": True, "size": 15, "space_after": 8},
        {"text": "C4.5 i Naive‑Bayes, 7 zbiorów UCI, 50 powtórzeń", "bullet": True, "size": 15, "space_after": 8},
        {"text": "CV pesymistyczna dla k = 2 i 5; k = 10–20 prawie bez obciążenia", "bullet": True, "size": 15, "space_after": 8},
        {"text": "Największa wariancja dla k = 2; stratyfikacja pomaga", "bullet": True, "size": 15, "space_after": 8},
        {"text": "Zalecenie: 10‑krotna stratyfikowana CV", "bullet": True, "size": 15, "bold": True},
    ])
    add_text(s, L + 6.3, TOP, 5.2, 0.4, [{"text": "Odtwarzamy: Rys. 1 i Rys. 3 (obciążenie i rozrzut vs k)", "bold": True, "size": 14, "color": NAVY}])
    rows = [
        ["", "Kohavi", "My (adaptacja)"],
        ["Dane", "7 zbiorów UCI", "generator syntetyczny"],
        ["Modele", "C4.5, Naive‑Bayes", "CART, GaussianNB"],
        ["Odniesienie", "holdout, 50 powtórzeń", "test N = 10 000, 30 powtórzeń"],
        ["k", "2, 5, 10, 20, LOO, bootstrap", "2, 5, 10"],
        ["Bez zmian", "te same foldy dla obu modeli; wiele niezależnych powtórzeń", "tak samo"],
    ]
    add_table(s, rows, L + 6.3, TOP + 0.45, 5.2, [1.2, 2.2, 1.8], row_h=0.5, size=11.5)
    add_notes(s, (
        "CO POWIEDZIEĆ (1,5 min): Kohavi porównał metody szacowania trafności klasyfikatora na małej próbce: holdout, "
        "k-krotną walidację krzyżową dla różnych k, leave-one-out i bootstrap. Zrobił to na dwóch różnych algorytmach "
        "(drzewo C4.5 i Naive-Bayes) i siedmiu zbiorach UCI, po 50 powtórzeń; „prawdziwą” trafność liczył holdoutem na "
        "reszcie zbioru. Wyniki: CV zaniża trafność (jest pesymistyczna), najbardziej dla 2 i 5 foldów; przy 10–20 "
        "foldach obciążenie prawie znika. Wariancja jest największa przy 2 foldach. Stratyfikacja zmniejsza i obciążenie, "
        "i wariancję. Bootstrap ma małą wariancję, ale na niektórych zbiorach ogromne obciążenie. Zalecenie: 10-krotna "
        "stratyfikowana CV.\n"
        "FRAGMENT DO ODTWORZENIA: Rysunek 1 (obciążenie CV w funkcji k) i Rysunek 3 (odchylenie standardowe w funkcji k) "
        "— dla naszych dwóch modeli i k = 2, 5, 10. Bootstrapu, leave-k-out i k = 20 nie odtwarzamy — karta tematu "
        "tego nie wymaga.\n"
        "ADAPTACJA (powiedzieć wprost): dane i modele są inne — narzuca je karta tematu. Zachowujemy cel i schemat "
        "badania: te same foldy dla obu modeli, wiele niezależnych powtórzeń, punkt odniesienia niezależny od CV.\n"
        "MOŻLIWE PYTANIE: „czym różni się Wasz punkt odniesienia od Kohaviego?” — on nie znał prawdy, więc szacował ją "
        "holdoutem na reszcie zbioru; my generujemy dane, więc mamy duży niezależny test 10 000 obiektów."
    ))

    # 3. Pytanie i hipotezy ------------------------------------------------------------------
    s = new_slide("Pytanie badawcze i hipotezy")
    add_card(s, L, TOP, W, 1.1, CARD_NAVY)
    add_text(s, L + 0.2, TOP + 0.08, W - 0.4, 0.95, [
        {"text": "Jak liczba części walidacji krzyżowej wpływa na obciążenie i zmienność szacowanego błędu oraz co zmienia stratyfikacja przy niezbalansowanych klasach?", "size": 17, "bold": True, "color": NAVY},
    ], anchor="middle")
    add_text(s, L, TOP + 1.35, W, 3.0, [
        {"runs": [("H1  ", {"bold": True, "color": NAVY}), ("Obciążenie pesymistyczne (CV zawyża błąd), największe dla k = 2, bliskie zera dla k = 10.", {})], "size": 17, "space_after": 14},
        {"runs": [("H1′ ", {"bold": True, "color": NAVY}), ("Zmienność największa dla k = 2.", {})], "size": 17, "space_after": 14},
        {"runs": [("H1″ ", {"bold": True, "color": NAVY}), ("Oba efekty silniejsze dla CART (niestabilny) niż dla GaussianNB (stabilny).", {})], "size": 17, "space_after": 14},
        {"runs": [("H2  ", {"bold": True, "color": NAVY}), ("Stratyfikacja przy klasie rzadkiej: mniejsza zmienność, brak foldów bez klasy rzadkiej, obciążenie podobne.", {})], "size": 17},
    ])
    add_notes(s, (
        "CO POWIEDZIEĆ (1 min): Przeczytać pytanie z karty tematu. Potem hipotezy z uzasadnieniem:\n"
        "H1 — model w foldzie uczy się na n(k−1)/k obiektach, czyli przy k = 2 tylko na 100 z 200. Mniej danych = "
        "większy błąd (krzywa uczenia), więc CV zawyża błąd; przy k = 10 trening ma 180 obiektów i różnica prawie znika.\n"
        "H1′ — przy k = 2 model jest „chwiejny”: inne 100 obiektów daje wyraźnie inny model, więc estymata skacze.\n"
        "H1″ — Kohavi pokazuje (Propozycja 1), że dla modelu stabilnego CV jest nieobciążona i jej wariancja nie zależy od k. "
        "GaussianNB uczy się tylko średnich i wariancji cech — jest stabilny; drzewo CART bez przycinania zmienia kształt "
        "przy każdej zmianie danych — jest niestabilne. Więc efekty powinny być widoczne głównie dla CART.\n"
        "H2 — stratyfikacja nie zmienia rozmiaru treningu (ta sama krzywa uczenia → obciążenie podobne), ale usuwa losowe "
        "wahania proporcji klas między foldami → mniejsza zmienność; i z definicji żaden fold nie zostaje bez klasy rzadkiej.\n"
        "PODKREŚLIĆ: wynik negatywny nie obniża oceny (kontrakt §6) — hipotezy to przewidywania, nie cel."
    ))

    # 4. Obciążenie i zmienność ------------------------------------------------------------------
    s = new_slide("Obciążenie i zmienność estymatora błędu")
    add_card(s, L, TOP, W, 1.2)
    add_text(s, L + 0.2, TOP + 0.08, 5.0, 0.4, [{"text": "Błąd prawdziwy modelu C (generalizacja)", "bold": True, "size": 15, "color": NAVY}])
    add_formula(s, r"e(C) = P\left(C(X) \neq Y\right)", L + 0.2, TOP + 0.55, 0.45, build_dir, "f_true", 26, center_in=5.0)
    add_text(s, L + 5.6, TOP + 0.05, W - 5.8, 1.1, [
        {"text": "Nieznany — przybliżamy go przez błąd eₜₑₛₜ: model nauczony na całym D (200 obiektów), zmierzony na niezależnym teście T (10 000 obiektów).", "size": 14, "color": INK2},
    ], anchor="middle")
    add_card(s, L, TOP + 1.35, W, 1.75, CARD_NAVY)
    add_text(s, L + 0.2, TOP + 1.43, 6.0, 0.4, [{"text": "Estymator CV i jego dwa rodzaje błędu", "bold": True, "size": 15, "color": NAVY}])
    add_formula(s, r"\mathrm{bias} = \mathrm{E}\left[\hat e_{CV} - e_{test}\right], \qquad \mathrm{sd}_{\Delta} = \mathrm{sd}\left(\hat e_{CV} - e_{test}\right)",
                L, TOP + 1.9, 0.5, build_dir, "f_bias_sd", 26, center_in=W)
    add_text(s, L + 0.2, TOP + 2.5, W - 0.4, 0.5, [
        {"text": "bias > 0 = pesymizm (CV zawyża błąd)      sd_Δ = odchylenie standardowe różnicy między powtórzeniami = zmienność", "size": 13, "color": INK2, "align": "center"},
    ], anchor="middle")
    add_text(s, L, TOP + 3.3, W, 1.1, [
        {"runs": [("Pesymizm: ", {"bold": True}), ("model w foldzie uczy się na n(k−1)/k obiektach — k = 2: 100, k = 5: 160, k = 10: 180 z 200.", {})], "size": 15, "bullet": True, "space_after": 7},
        {"runs": [("Powtórzenie: ", {"bold": True}), ("niezależny zbiór Dᵣ, r = 1…30 → 30 różnic; ich średnia = bias, odchylenie standardowe = zmienność.", {})], "size": 15, "bullet": True, "space_after": 7},
        {"runs": [("Niska zmienność ≠ małe obciążenie.", {"bold": True})], "size": 15, "bullet": True},
    ])
    add_notes(s, (
        "CO POWIEDZIEĆ (1,5 min): Generalizacja: błąd prawdziwy e(C) to prawdopodobieństwo pomyłki na NOWYM obiekcie "
        "z tego samego rozkładu. Nie da się go policzyć dokładnie — w praktyce nie mamy nowych danych. My mamy generator, "
        "więc tworzymy duży niezależny test T (10 000 obiektów) i błąd na nim, e_test, traktujemy jako prawdę (jego własna "
        "niepewność to ok. 0,4 punktu procentowego).\n"
        "Obciążenie estymatora = wartość oczekiwana różnicy między estymatą ê_CV a prawdą e_test; E[·] to średnia po "
        "nieskończenie wielu powtórzeniach — my przybliżamy ją średnią z 30. Dodatnie obciążenie = CV zawyża błąd = "
        "pesymizm (Kohavi mówi o zaniżonej trafności — to samo). Zmienność = odchylenie standardowe tej różnicy między "
        "powtórzeniami, czyli „jak bardzo wynik skacze”. Dokładnie tak definiuje to karta tematu.\n"
        "Trening/walidacja: w CV model widzi tylko część danych — przy k = 2 połowę. Mniej danych → większy błąd → "
        "ê_CV mierzy błąd gorszego modelu niż ten na całych 200 → pesymizm. Im większe k, tym większy trening w pojedynczym "
        "dopasowaniu i tym mniejszy pesymizm.\n"
        "Powtórzenie badania to CAŁY niezależny zbiór D_r (nie fold) — bo foldy nie są niezależne (następny slajd).\n"
        "PYTANIE KONTROLNE 2: „czy niższa wariancja oznacza mniejsze obciążenie?” — NIE, to dwie osobne rzeczy "
        "(strzelec: można strzelać powtarzalnie i zawsze obok). Przykład Kohaviego: bootstrap .632 ma małą wariancję "
        "i obciążenie do 9,8 p.p."
    ))

    # 5. Wzór CV i symbole -------------------------------------------------------------------
    s = new_slide("k‑krotna walidacja krzyżowa: wzór i symbole")
    add_formula(s, r"\hat e_{CV} = \frac{1}{k}\sum_{j=1}^{k} e_j, \qquad e_j = \frac{m_j}{|D_j|}",
                L, TOP, 0.85, build_dir, "f_cv", 30, center_in=W)
    symbols = [
        (r"D = D_1 \cup \dots \cup D_k", "zbiór treningowy (n = 200) podzielony losowo na k równych, rozłącznych foldów"),
        (r"C_{-j}", "model nauczony bez foldu j (na n(k−1)/k obiektach), sprawdzany na odłożonym foldzie Dⱼ"),
        (r"m_j,\ e_j", "liczba pomyłek modelu C₋ⱼ na foldzie Dⱼ i ich odsetek (pomyłki / rozmiar foldu)"),
        (r"\hat e_{CV}", "średnia błędów foldów; przy równych foldach = wszystkie pomyłki / n"),
        (r"e_{test}", "błąd modelu nauczonego na całym D, zmierzony na teście T (N = 10 000)"),
    ]
    row_y, row_h = TOP + 1.05, 0.5
    for i, (tex, desc) in enumerate(symbols):
        y = row_y + i * row_h
        add_card(s, L, y, W, row_h - 0.06, CARD if i % 2 == 0 else CARD_NAVY)
        add_formula(s, tex, L + 0.15, y + 0.09, 0.3, build_dir, f"sym_{i}", 22, center_in=2.4)
        add_text(s, L + 2.7, y, W - 2.9, row_h - 0.06, [{"text": desc, "size": 14, "color": INK}], anchor="middle")
    add_text(s, L, row_y + len(symbols) * row_h + 0.1, W, 0.6, [
        {"runs": [("Foldy nie są niezależne: ", {"bold": True}), ("zbiory uczące C₋₁ i C₋₂ mają wspólne (k−2)/(k−1) obiektów (k = 10: 8 z 9).", {})], "size": 14.5},
    ])
    add_notes(s, (
        "CO POWIEDZIEĆ (1,5 min): Przeczytać wzór symbol po symbolu. Zbiór D dzielimy losowo na k foldów. Dla każdego "
        "foldu j uczymy model na pozostałych k−1 foldach (C_−j, „minus j” = bez j-tego) i sprawdzamy na odłożonym D_j: "
        "m_j pomyłek na |D_j| obiektów daje e_j. Estymata ê_CV to zwykła średnia z k błędów foldów. Daszek = oszacowanie, "
        "nie prawda. Każdy obiekt jest dokładnie raz w foldzie testowym, więc przy równych foldach średnia = (wszystkie "
        "pomyłki)/n — to wzór (4) Kohaviego; u nas 200 dzieli się przez 2, 5 i 10, więc obie wersje są identyczne.\n"
        "Koszt: k dopasowań modelu na jedno ê_CV.\n"
        "PYTANIE KONTROLNE 1: „dlaczego 10 foldów nie daje 10 niezależnych pomiarów?” — zbiory uczące sąsiednich foldów "
        "pokrywają się w 8 z 9 (ogólnie (k−2)/(k−1)); to prawie te same modele, więc e_j są skorelowane. Odchylenie "
        "między foldami podzielone przez pierwiastek z k zaniża niepewność. Skrajny przykład Kohaviego: iris + model "
        "„zawsze klasa większościowa” + leave-one-out → każdy fold 0 % trafności, rozrzut 0, a prawda 33 %. Dlatego "
        "jednostką powtórzenia jest cały niezależny zbiór D_r."
    ))

    # 6. Przykład ręczny ---------------------------------------------------------------------
    ex = folds[(folds.rep == 1) & (folds.model == "gnb") & (folds.k == 5) & (folds.cv_type == "kfold")].sort_values("fold")
    ex_sum = summary[(summary.rep == 1) & (summary.model == "gnb") & (summary.k == 5)].iloc[0]
    s = new_slide("Przykład ręczny: średnia błędów foldów (D₁, GaussianNB, k = 5)")
    rows = [["Fold j", "Obiektów w foldzie", "Pomyłek mⱼ", "eⱼ = mⱼ / 40"]]
    for _, row in ex.iterrows():
        rows.append([str(int(row.fold)), str(int(row.n_test)), str(int(row.n_errors)), f"{int(row.n_errors)} / {int(row.n_test)} = {num(row.error)}"])
    total_err, total_n = int(ex.n_errors.sum()), int(ex.n_test.sum())
    rows.append(["Razem", str(total_n), str(total_err), f"{total_err} / {total_n} = {num(total_err / total_n)}"])
    add_table(s, rows, L, TOP, 6.4, [1.0, 1.8, 1.3, 2.3], row_h=0.42, size=14, align=["c", "c", "c", "c"], bold_last=True)
    terms = " + ".join(tex_num(e) for e in ex.error)
    add_card(s, L + 6.7, TOP, 4.8, 2.95, CARD_NAVY)
    add_text(s, L + 6.9, TOP + 0.1, 4.4, 0.4, [{"text": "Rachunek", "bold": True, "size": 16, "color": NAVY}])
    add_formula(s, rf"\hat e_{{CV}} = \frac{{1}}{{5}}\,({terms})", L + 6.85, TOP + 0.5, 0.45, build_dir, "f_ex1", 22, center_in=4.5)
    add_formula(s, rf"= {tex_num(ex.error.sum())} / 5 = \mathbf{{{tex_num(ex_sum.e_cv)}}}", L + 6.85, TOP + 1.0, 0.34, build_dir, "f_ex1b", 22, center_in=4.5)
    add_formula(s, rf"e_{{test}} = \mathbf{{{tex_num(ex_sum.e_test, 4)}}}", L + 6.85, TOP + 1.55, 0.34, build_dir, "f_ex2", 22, center_in=4.5)
    add_text(s, L + 6.9, TOP + 1.92, 4.4, 0.3, [{"text": "(model na całym D₁, test 10 000 obiektów)", "size": 12, "color": INK2, "align": "center"}])
    add_formula(s, rf"\hat e_{{CV}} - e_{{test}} = \mathbf{{{tex_num(ex_sum['diff'], 4)}}}", L + 6.85, TOP + 2.35, 0.38, build_dir, "f_ex3", 22, center_in=4.5)
    add_text(s, L, TOP + 3.1, W, 1.0, [
        {"text": "Jedna różnica to przypadek — obciążenie i zmienność liczymy z 30 takich różnic (po jednej na zbiór Dᵣ).", "size": 15, "color": INK2, "bullet": True},
        {"text": f"Liczby z próby uruchomienia (results/pilot, {meta['timestamp'][:10]}).", "size": 15, "color": INK2, "bullet": True},
    ])
    add_notes(s, (
        "CO POWIEDZIEĆ (1,5 min) — POLICZYĆ NA GŁOS: zbiór D_1 z naszego generatora, 200 obiektów, GaussianNB, 5 foldów "
        "po 40. Pomyłki na foldach: 7, 3, 7, 8, 4. Błędy: 7/40 = 0,175; 3/40 = 0,075; 7/40 = 0,175; 8/40 = 0,200; "
        "4/40 = 0,100. Suma 0,725, podzielić przez 5: ê_CV = 0,145. Kontrola drugim sposobem: 29 pomyłek / 200 obiektów "
        "= 0,145 — to samo, bo foldy są równe.\n"
        "Punkt odniesienia: ten sam GaussianNB nauczony na całych 200 obiektach, sprawdzony na teście T: 1754 pomyłki "
        "na 10 000 → e_test = 0,1754. Różnica = 0,145 − 0,1754 = −0,030: w tym konkretnym zbiorze CV ZANIŻYŁA błąd "
        "o 3 p.p. — czyli optymizm, wbrew hipotezie. To nie przeczy H1: obciążenie to ŚREDNIA z wielu zbiorów, a "
        "pojedynczy błąd foldu z 40 obiektów ma rozrzut ok. 6 p.p. Dopiero 30 różnic (po jednej na niezależny zbiór) "
        "da średnią (obciążenie) i odchylenie standardowe (zmienność).\n"
        "JEŚLI ZAPYTAJĄ O SD: na 3 liczbach, np. różnice 0,01; 0,03; 0,05 → średnia 0,03, odchylenia −0,02/0/+0,02, "
        "kwadraty 0,0004/0/0,0004, suma 0,0008 / (3−1) = 0,0004, pierwiastek 0,02."
    ))

    # 7. Generator, powtórzenia, ocena referencyjna --------------------------------------------
    s = new_slide("Generator danych, powtórzenia i ocena referencyjna")
    add_formula(s, r"X \mid y \sim \mathcal{N}(\mu_y, I_{10}),\quad \mu_1 = -\mu_0 = \frac{\delta}{2}\,u,\quad u = \frac{1}{\sqrt{5}}(1,1,1,1,1,0,0,0,0,0),\quad \delta = 2,\quad y \sim \mathrm{Bernoulli}(\pi_1)",
                L, TOP - 0.05, 0.5, build_dir, "f_gen", 28, center_in=W)
    img_h = 2.7
    img_w = img_h * 12.0 / 4.2    # proporcje rysunku z make_figures.py
    s.shapes.add_picture(str(figures / "generator.png"), Inches(L + (W - img_w) / 2), Inches(TOP + 0.55), height=Inches(img_h))
    add_text(s, L, TOP + 0.55 + img_h + 0.1, W, 1.0, [
        {"runs": [("Generator: ", {"bold": True}), (f"dwie klasy normalne, 5 cech informacyjnych + 5 szumu; błąd Bayesa {pct(gen.bayes_error())} (π₁ = 0,5), {pct(gen_imb.bayes_error())} (π₁ = 0,1).", {})], "size": 14.5, "bullet": True, "space_after": 5},
        {"runs": [("Powtórzenia: ", {"bold": True}), ("R = 30 zbiorów Dᵣ po n = 200.   ", {}), ("Ocena referencyjna: ", {"bold": True}), ("jeden test T, N = 10 000, osobne ziarno, tylko do oceny.", {})], "size": 14.5, "bullet": True},
    ])
    add_notes(s, (
        "CO POWIEDZIEĆ (1,5 min): Karta tematu każe ustalić generator, liczbę powtórzeń i ocenę referencyjną.\n"
        "GENERATOR: każdy obiekt ma 10 cech i klasę y. Najpierw losujemy klasę (Bernoulli: π₁ = 0,5 w odtworzeniu, 0,1 "
        "w eksperymencie własnym), potem 10 cech z rozkładu normalnego o środku zależnym od klasy. Środki μ₀ i μ₁ leżą "
        "symetrycznie w odległości δ = 2 wzdłuż kierunku u: pierwsze 5 cech różni klasy (informacyjne), ostatnie 5 to "
        "czysty szum. Rysunek: po lewej 200 wylosowanych punktów na dwóch cechach — klasy się mieszają; po prawej oba "
        "„dzwony” po rzucie na kierunek u. Zacieniowane pola = błąd Bayesa, czyli błąd najlepszego możliwego klasyfikatora: "
        "Φ(−δ/2) = Φ(−1) = 15,9 %. Przy klasie rzadkiej granica przesuwa się w stronę klasy rzadkiej (próg ln 9/δ ≈ 1,10) "
        "i błąd Bayesa to 7,0 % — ale „zawsze klasa 0” ma 10 %, dlatego w eksperymencie własnym mierzymy też czułość.\n"
        "DLACZEGO TAKI: znamy błąd Bayesa, więc wiemy, jak daleko od ideału są modele. GaussianNB zakłada dokładnie ten "
        "rozkład (niezależne cechy normalne) → stabilny; CART tnie równolegle do osi, a granica jest ukośna → musi ją "
        "schodkować, przeucza się → niestabilny. To odpowiednik pary Naive-Bayes / C4.5 u Kohaviego.\n"
        "POWTÓRZENIA: 30 niezależnych zbiorów po 200 (z karty), ziarno 1000 + r. OCENA REFERENCYJNA: jeden test T, "
        "10 000 obiektów, osobne ziarno; służy WYŁĄCZNIE do oceny — nic według niego nie dobieramy (pytanie kontrolne 3). "
        "Parametry generatora ustaliliśmy z teorii (błąd Bayesa ok. 16 %), zanim cokolwiek policzyliśmy.\n"
        "ZMIANA PARAMETRU (gdyby pytali): większe δ → mniejszy błąd Bayesa, płaska krzywa uczenia, mniejszy pesymizm; "
        "więcej cech szumowych → większy błąd i pesymizm, zwłaszcza dla CART."
    ))

    # 8. Plan odtworzenia (schemat) ---------------------------------------------------------
    s = new_slide("Plan odtworzenia: estymata CV kontra duży test")
    bx_h = 0.85
    row1_y, row2_y = TOP, TOP + 1.4

    def box(x, y, w, h, head, body, fill=CARD, head_tex=None):
        add_card(s, x, y, w, h, fill)
        if head_tex:
            add_formula(s, head_tex, x + 0.12, y + 0.08, 0.28, build_dir, f"box_{head_tex[:6]}_{int(x*10)}", 20, color="#1f3a5f")
            add_text(s, x + 0.12, y + 0.42, w - 0.24, h - 0.46, [{"text": body, "size": 11.5, "color": INK2}])
        else:
            add_text(s, x + 0.12, y + 0.06, w - 0.24, h - 0.1, [
                {"text": head, "bold": True, "size": 13, "color": NAVY, "space_after": 2},
                {"text": body, "size": 11.5, "color": INK2},
            ])

    box(L, row1_y, 2.2, bx_h, "Generator", "ziarno 1000 + r dla Dᵣ; 90001 dla T")
    add_arrow(s, L + 2.25, row1_y + 0.28, 0.35, 0.3)
    box(L + 2.65, row1_y, 2.3, bx_h, None, "n = 200, klasy 50/50", CARD_NAVY, head_tex=r"D_r,\ r = 1,\dots,30")
    add_arrow(s, L + 5.0, row1_y + 0.28, 0.35, 0.3)
    box(L + 5.4, row1_y, 3.1, bx_h, "CV, k ∈ {2, 5, 10}", "KFold(shuffle, ziarno 2000 + r); te same podziały dla obu modeli")
    add_arrow(s, L + 8.55, row1_y + 0.28, 0.35, 0.3)
    box(L + 8.95, row1_y, 2.55, bx_h, None, "średnia błędów k foldów", CARD_NAVY, head_tex=r"\hat e_{CV,r,k}")
    add_arrow(s, L + 3.65, row1_y + bx_h + 0.04, 0.3, 0.46, "down")
    box(L + 2.65, row2_y, 2.3, bx_h, "Model na całym Dᵣ", "jedno dopasowanie, 200 obiektów")
    add_arrow(s, L + 5.0, row2_y + 0.28, 0.35, 0.3)
    box(L + 5.4, row2_y, 3.1, bx_h, "Test T", "N = 10 000; tylko do oceny")
    add_arrow(s, L + 8.55, row2_y + 0.28, 0.35, 0.3)
    box(L + 8.95, row2_y, 2.55, bx_h, None, "odsetek pomyłek na T", CARD_NAVY, head_tex=r"e_{test,r}")
    add_arrow(s, L + 10.1, row2_y + bx_h + 0.05, 0.3, 0.38, "down")
    add_card(s, L, row2_y + bx_h + 0.48, W, 0.8, CARD_NAVY)
    add_formula(s, r"\Delta_r = \hat e_{CV,r} - e_{test,r}", L + 0.25, row2_y + bx_h + 0.64, 0.42, build_dir, "f_delta", 24)
    add_text(s, L + 4.2, row2_y + bx_h + 0.5, W - 4.4, 0.76, [
        {"text": "→ 30 wartości dla każdego k i modelu: średnia = obciążenie, odchylenie standardowe = zmienność", "size": 14.5},
    ], anchor="middle")
    add_text(s, L, row2_y + bx_h + 1.45, W, 0.6, [
        {"runs": [("Modele: ", {"bold": True}), ("GaussianNB, CART (DecisionTreeClassifier, domyślne).   ", {}), ("Metryka: ", {"bold": True}), ("odsetek pomyłek.   ", {}), ("Utrwalone w: ", {"bold": True}), ("cvbias/protocol.py", {})], "size": 14, "color": INK2},
    ])
    add_notes(s, (
        "CO POWIEDZIEĆ (1 min) — przejść po strzałkach: generator tworzy 30 niezależnych zbiorów D_r po 200 obiektów "
        "(ziarno 1000 + r) i jeden test T z osobnego ziarna. Dla każdego zbioru i każdego z dwóch modeli: górna ścieżka — "
        "CV z k = 2, 5, 10 na tym samym zbiorze i z tymi samymi podziałami dla obu modeli (jak u Kohaviego: ziarno "
        "podziału zależy tylko od r) → trzy estymaty ê_CV; dolna ścieżka — model na całym D_r sprawdzony na T → e_test. "
        "Różnica Δ = ê_CV − e_test; dla każdego k i modelu mamy 30 niezależnych Δ → średnia to obciążenie, odchylenie "
        "standardowe to zmienność (dokładnie wielkości z karty tematu).\n"
        "Wszystko utrwalone w cvbias/protocol.py przed badaniem: rozmiary, ziarna, k, modele, metryka — karta: „przed "
        "końcowym badaniem utrwal podziały danych, parametry, metryki i losowania”.\n"
        "Koszt: na zbiór i model 2 + 5 + 10 dopasowań w CV + 1 na całości = 18; razem 30 · 2 · 18 = 1080.\n"
        "GDZIE W KODZIE (gdyby prosili): experiment.run_series — generate(T) raz przed pętlą; pętla po r; "
        "fit_and_predict na całym D_r → e_test; make_splitter(…, seeds['split'] + r) → te same podziały; "
        "cross_validate → błędy foldów."
    ))

    # 9. Eksperyment własny ----------------------------------------------------------------
    s = new_slide("Eksperyment własny: stratyfikacja przy klasie rzadkiej (≈ 10 %)")
    add_card(s, L, TOP, 5.6, 4.3)
    add_text(s, L + 0.2, TOP + 0.15, 5.2, 4.1, [
        {"text": "Plan", "bold": True, "size": 16, "color": NAVY, "space_after": 10},
        {"text": "Generator z π₁ = 0,1 → ok. 20 obiektów klasy rzadkiej na 200; 30 zbiorów, osobny test T′.", "size": 15, "bullet": True, "space_after": 10},
        {"text": "Na tym samym Dᵣ: KFold(5) i StratifiedKFold(5), to samo ziarno.", "size": 15, "bullet": True, "space_after": 10},
        {"text": "Miary: odsetek pomyłek, czułość klasy rzadkiej, liczba foldów bez klasy rzadkiej.", "size": 15, "bullet": True, "space_after": 10},
        {"text": "Fold bez klasy rzadkiej → czułość nieokreślona → zapis jako brak wyniku z powodem.", "size": 15, "bullet": True, "bold": True},
    ])
    own1 = own_folds[(own_folds.rep == 1) & (own_folds.model == "gnb")]
    kf = own1[own1.cv_type == "kfold"].sort_values("fold")
    st = own1[own1.cv_type == "stratified"].sort_values("fold")
    rows = [["Fold", "KFold: klasa rzadka", "Stratified: klasa rzadka"]]
    for (_, a), (_, b) in zip(kf.iterrows(), st.iterrows()):
        rows.append([str(int(a.fold)), str(int(a.n_test_class1)), str(int(b.n_test_class1))])
    n_rare_d1 = int(own1.n_train_class1.iloc[0] + own1.n_test_class1.iloc[0])
    add_text(s, L + 5.9, TOP, 5.6, 0.5, [{"text": f"Próba: D₁ ({n_rare_d1} obiektów klasy rzadkiej na 200), foldy po 40", "size": 13, "color": INK2}])
    add_table(s, rows, L + 5.9, TOP + 0.5, 5.6, [0.8, 2.4, 2.4], row_h=0.42, size=14, align=["c", "c", "c"])
    n_missing = int(own_folds[own_folds.cv_type == "kfold"].missing_class_test.sum())
    n_missing_s = int(own_folds[own_folds.cv_type == "stratified"].missing_class_test.sum())
    add_text(s, L + 5.9, TOP + 3.1, 5.6, 1.1, [
        {"text": f"W 3 zbiorach próby: foldy bez klasy rzadkiej — KFold {n_missing}, Stratified {n_missing_s}.", "size": 14, "color": INK2, "bullet": True},
    ])
    add_notes(s, (
        "CO POWIEDZIEĆ (1 min): Karta: udział klasy rzadkiej ok. 10 %, porównanie zwykłej i stratyfikowanej 5-foldowej "
        "CV NA TYCH SAMYCH PRÓBKACH. Stratyfikacja = foldy mają proporcje klas takie jak cały zbiór (StratifiedKFold losuje "
        "osobno w każdej klasie). Bierzemy ten sam zbiór D_r i dzielimy go dwoma sposobami z tym samym ziarnem — różni się "
        "tylko przydział do foldów.\n"
        "Tabela z próby: w D_1 jest 24 obiekty klasy rzadkiej; zwykły KFold rozrzucił je 2/5/9/6/2, stratyfikowany 4/5/5/5/5. "
        "W zbiorze D_3 (14 rzadkich) zwykły podział dał fold bez ani jednego obiektu klasy rzadkiej — wtedy czułość to 0/0, "
        "nieokreślona. Karta każe takie przypadki zapisywać: w CSV jest flaga i powód; nie wstawiamy zera, nie pomijamy "
        "po cichu, a liczba takich foldów jest jedną z miar.\n"
        "Dlaczego czułość: przy 10 % klasy rzadkiej model „zawsze 0” ma tylko 10 % błędu — odsetek pomyłek nie pokazuje, "
        "czy klasa rzadka jest w ogóle wykrywana.\n"
        "Oczekiwanie (H2): obciążenie podobne (ten sam rozmiar treningu w foldzie), zmienność mniejsza (znika losowość "
        "proporcji klas między foldami) — zgodne z obserwacją Kohaviego, że stratyfikacja jest „jednolicie lepsza”."
    ))

    # 10. Wykonalność i kontrola poprawności ------------------------------------------------
    s = new_slide("Wykonalność: próba uruchomienia i kontrola poprawności")
    tiles = [
        (str(meta["fits_performed"]), "dopasowań w próbie", "3 zbiory, oba eksperymenty"),
        (f"{num(meta['elapsed_seconds'], 2)} s", "czas próby na CPU", "zwykły laptop, bez GPU"),
        (f"{fits['total']} / {fits['budget']}", "dopasowań w pełnej serii", "budżet z karty tematu"),
    ]
    for i, (big, label, sub) in enumerate(tiles):
        x = L + i * 3.9
        add_card(s, x, TOP, 3.7, 1.4, CARD_NAVY)
        add_text(s, x + 0.15, TOP + 0.08, 3.4, 0.7, [{"text": big, "size": 32, "bold": True, "color": NAVY, "align": "center"}], anchor="middle")
        add_text(s, x + 0.15, TOP + 0.78, 3.4, 0.6, [
            {"text": label, "size": 12, "color": INK2, "align": "center", "space_after": 0},
            {"text": sub, "size": 12, "color": MUTED, "align": "center", "space_after": 0},
        ])
    add_text(s, L, TOP + 1.65, 5.6, 2.6, [
        {"text": "Środowisko", "bold": True, "size": 15, "color": NAVY, "space_after": 6},
        {"text": f"Python {meta['versions']['python']}, scikit‑learn {meta['versions']['scikit-learn']}", "size": 15, "bullet": True, "space_after": 6},
        {"text": "Zapisujemy błąd i skład klas każdego foldu, podsumowania, meta.json (wersje, ziarna, liczba dopasowań)", "size": 15, "bullet": True, "space_after": 6},
        {"text": "github.com/FylyppS/rozpoznawanie-wzorcow", "size": 15, "bullet": True},
    ])
    add_text(s, L + 5.9, TOP + 1.65, 5.6, 2.6, [
        {"text": "Kontrola poprawności", "bold": True, "size": 15, "color": NAVY, "space_after": 6},
        {"text": "Duży test tylko do oceny — nic według niego nie dobieramy", "size": 15, "bullet": True, "space_after": 6},
        {"text": "Te same zbiory i podziały dla wszystkich wariantów CV i obu modeli", "size": 15, "bullet": True, "space_after": 6},
        {"text": "Ziarna, rozmiary, k, metryki utrwalone przed badaniem", "size": 15, "bullet": True, "space_after": 6},
        {"text": "Brak przecieku: model w foldzie widzi tylko D ∖ Dⱼ", "size": 15, "bullet": True},
    ])
    add_notes(s, (
        "CO POWIEDZIEĆ (1 min): Kontrakt wymaga na etapie 1 „małej próby uruchomienia”. Uruchomiliśmy cały potok na "
        "3 zbiorach zamiast 30: 174 dopasowania w 0,16 s — pełna seria to 1080 (odtworzenie) + 660 (eksperyment własny) "
        "+ 180 (etap 3: modele na 100/160/180 obiektach, żeby oddzielić wpływ rozmiaru treningu od losowego podziału) "
        "= 1920 ≤ 2500. Obliczenia nie są ryzykiem; ryzykiem jest poprawność protokołu.\n"
        "Z próby NIE wyciągamy wniosków — służy do sprawdzenia kodu i do przykładu ręcznego.\n"
        "Co zapisujemy (karta: „zachowaj wyniki wszystkich zaplanowanych uruchomień, błędy oraz powody brakujących "
        "wyników”): CSV per fold z błędem i składem klas, podsumowania per zbiór × model × k, meta.json z wersjami, "
        "ziarnami, liczbą dopasowań i czasem.\n"
        "Kontrola poprawności — cztery zasady z karty: (1) duży test wyłącznie do oceny estymatora; (2) te same próbki "
        "dla wszystkich wariantów CV; (3) wszystko utrwalone przed badaniem w protocol.py; (4) brak przecieku — model "
        "w foldzie uczy się tylko na pozostałych foldach, T ma osobne ziarno, nie ma doboru hiperparametrów ani "
        "skalowania, więc nie ma kroku, który mógłby „podejrzeć” dane testowe."
    ))

    # 11. Ograniczenia i źródła --------------------------------------------------------------
    s = new_slide("Ograniczenia planu i źródła")
    add_card(s, L, TOP, 5.6, 4.0)
    add_text(s, L + 0.2, TOP + 0.15, 5.2, 3.8, [
        {"text": "Ograniczenia", "bold": True, "size": 16, "color": NAVY, "space_after": 10},
        {"text": "Dane „idealne” dla GaussianNB (rozkład zgodny z jego założeniami).", "size": 15, "bullet": True, "space_after": 10},
        {"text": "Jeden wspólny test T dla wszystkich powtórzeń.", "size": 15, "bullet": True, "space_after": 10},
        {"text": "R = 30: efekty poniżej ok. 1 p.p. nieodróżnialne od zera.", "size": 15, "bullet": True, "space_after": 10},
        {"text": "Kohavi: „prawda” z holdoutu; próbki z jednego zbioru; Propozycja 1 zakłada stabilny model.", "size": 15, "bullet": True},
    ])
    add_card(s, L + 5.9, TOP, 5.6, 4.0, CARD_NAVY)
    add_text(s, L + 6.1, TOP + 0.15, 5.2, 3.8, [
        {"text": "Źródła", "bold": True, "size": 16, "color": NAVY, "space_after": 10},
        {"text": "R. Kohavi, A Study of Cross-Validation and Bootstrap for Accuracy Estimation and Model Selection, IJCAI 1995 — ai.stanford.edu/~ronnyk/accEst.pdf", "size": 14, "bullet": True, "space_after": 10},
        {"text": "scikit-learn 1.8.0: KFold, StratifiedKFold, GaussianNB, DecisionTreeClassifier", "size": 14, "bullet": True, "space_after": 10},
        {"text": "Karta tematu 06 i kontrakt przedmiotu (Moodle PO)", "size": 14, "bullet": True, "space_after": 10},
        {"text": "Kod, protokół i surowe wyniki: github.com/FylyppS/rozpoznawanie-wzorcow", "size": 14, "bullet": True},
    ])
    add_notes(s, (
        "CO POWIEDZIEĆ (1 min): Ograniczenia naszego planu: (1) dane syntetyczne o rozkładzie dokładnie zgodnym "
        "z założeniami GaussianNB — wyniki dla NB mogą być „zbyt ładne”; (2) jeden wspólny test T — wartości e_test "
        "z różnych zbiorów dzielą ten sam szum testu (ok. 0,4 p.p., pomijalny, ale odnotowany); (3) 30 powtórzeń → błąd "
        "standardowy obciążenia ≈ sd/√30, przy sd 3 p.p. to ≈ 0,5 p.p., więc obciążenie poniżej 1 p.p. (prawdopodobne "
        "dla NB) może być nieodróżnialne od zera — to też jest wynik, nie porażka.\n"
        "Krytycznie o publikacji: „prawdziwa” trafność Kohaviego to też tylko estymata z holdoutu na reszcie zbioru; "
        "50 próbek z tego samego skończonego zbioru nie jest w pełni niezależnych; jego Propozycja 1 (wariancja CV "
        "niezależna od k) wymaga stabilnego modelu, a drzewa tego nie spełniają — nasz generator daje prawdziwie "
        "niezależne zbiory i prawdziwy, duży test.\n"
        "Źródła: publikacja, dokumentacja scikit-learn, karta tematu i kontrakt. Praca autora vs nasz wkład: od Kohaviego "
        "pytanie, definicje i schemat porównania; generator, protokół, kod i eksperyment własny są nasze.\n"
        "Informacja o AI (załącznik 2A) jest osobnym PDF-em dołączonym do zadania w Moodle.\n"
        "NA PYTANIA — trzy pytania kontrolne z karty: (1) 10 foldów ≠ 10 niezależnych pomiarów, bo zbiory uczące "
        "pokrywają się w 8/9; (2) niższa wariancja ≠ mniejsze obciążenie (bootstrap .632: mała wariancja, obciążenie "
        "do 9,8 p.p.); (3) według dużego testu nie wolno dobierać parametrów modeli, generatora ani wariantu CV."
    ))

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    print(f"zapisano {out} ({len(prs.slides)} slajdów)")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--source-template", type=Path, default=None, help="pełna prezentacja z formatką (np. ~/Downloads/preznainz.pptx)")
    parser.add_argument("--template", type=Path, default=ROOT / "etap1/formatka-wi.pptx")
    parser.add_argument("--out", type=Path, default=ROOT / "etap1/slajdy-etap1.pptx")
    parser.add_argument("--results", type=Path, default=ROOT / "results/pilot")
    parser.add_argument("--figures", type=Path, default=ROOT / "etap1/figures")
    parser.add_argument("--build-dir", type=Path, default=ROOT / "build/slides")
    args = parser.parse_args()

    args.build_dir.mkdir(parents=True, exist_ok=True)
    if args.source_template is not None or not args.template.exists():
        source = args.source_template or Path.home() / "Downloads/preznainz.pptx"
        make_blank_template(source, args.template)
        print(f"formatka bez slajdów: {args.template}")
    build(args.template, args.out, args.results, args.figures, args.build_dir)


if __name__ == "__main__":
    main()
